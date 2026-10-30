"""Voter roll import, from a CSV or an Excel (.xlsx) sheet.

The admin uploads a file whose first row has these headers (any order, any case;
extra columns are ignored):

    full_name, reg_number, faculty, department, level

Existing students (matched by reg_number) are updated; their face
registration is kept. New students are added.
"""

import csv
import io
import zipfile
from dataclasses import dataclass, field

from django.db import transaction

from .models import Voter

REQUIRED_HEADERS = ["full_name", "reg_number", "faculty", "department", "level"]


def normalize(value):
    """Collapse whitespace and ignore case, for cross-checking typed details."""
    return " ".join((value or "").split()).casefold()


def clean_reg_number(value):
    return "".join((value or "").split()).upper()


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    errors: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


def _cell(value):
    """Excel stores 200 as 200.0; turn cells back into the text the admin typed."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def read_csv(uploaded_file):
    try:
        text = uploaded_file.read().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ValueError("The file is not UTF-8 text. Save it from Excel as 'CSV UTF-8', or upload the .xlsx file.")
    rows = list(csv.reader(io.StringIO(text)))
    return (rows[0] if rows else []), rows[1:]


def read_xlsx(uploaded_file):
    from openpyxl import load_workbook

    try:
        sheet = load_workbook(uploaded_file, read_only=True, data_only=True).worksheets[0]
    except (zipfile.BadZipFile, KeyError, OSError, IndexError):
        raise ValueError("This Excel file could not be read. Open it in Excel and save it again as .xlsx.")
    rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    return (rows[0] if rows else []), rows[1:]


def import_roll(uploaded_file):
    """Validate the whole file first; write nothing if any row is bad."""
    result = ImportResult()
    reader = read_xlsx if uploaded_file.name.lower().endswith(".xlsx") else read_csv
    try:
        header_row, data_rows = reader(uploaded_file)
    except ValueError as exc:
        result.errors.append(str(exc))
        return result

    headers = [normalize(h).replace(" ", "_") for h in header_row]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        result.errors.append(
            f"Missing column(s): {', '.join(missing)}. Required headers: {', '.join(REQUIRED_HEADERS)}."
        )
        return result

    rows, seen = [], {}
    for line_no, values in enumerate(data_rows, start=2):
        raw = dict(zip(headers, values))
        row = {h: " ".join((raw.get(h) or "").split()) for h in REQUIRED_HEADERS}
        if not any(row.values()):
            continue
        row["reg_number"] = clean_reg_number(row["reg_number"])
        empty = [h for h in REQUIRED_HEADERS if not row[h]]
        if empty:
            result.errors.append(f"Row {line_no}: empty {', '.join(empty)}.")
            continue
        if row["reg_number"] in seen:
            result.errors.append(
                f"Row {line_no}: reg number {row['reg_number']} already appears on row {seen[row['reg_number']]}."
            )
            continue
        seen[row["reg_number"]] = line_no
        rows.append(row)

    if not rows and not result.errors:
        result.errors.append("The file has headers but no students.")
    if result.errors:
        return result

    with transaction.atomic():
        existing = {v.reg_number: v for v in Voter.objects.filter(reg_number__in=seen)}
        for row in rows:
            voter = existing.get(row["reg_number"])
            if voter is None:
                Voter.objects.create(**row)
                result.created += 1
            else:
                for key in ("full_name", "faculty", "department", "level"):
                    setattr(voter, key, row[key])
                voter.save()
                result.updated += 1
    return result
